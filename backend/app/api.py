from __future__ import annotations

import csv
import io
import re
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from os.path import exists
from typing import Literal
from uuid import UUID

import jwt
from fastapi import APIRouter, Depends, Header, HTTPException, Query, Response, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from app.audit import audit, notify
from app.billing import (
    create_initial_period,
    current_period,
    ensure_period,
    extend_period,
    mark_paid_and_advance,
    period_balance,
    period_paid,
    set_period_dates,
)
from app.config import settings
from app.db import get_db
from app.dependencies import get_current_user, require_billing_permission, require_operator
from app.models import (
    AuditEvent,
    ApplicationFile,
    BillingChange,
    BillingPeriod,
    Dormitory,
    Expense,
    Notification,
    Payment,
    PaymentAdjustment,
    Resident,
    ResidencePeriod,
    Room,
    SystemSetting,
    StudentApplication,
    User,
)
from app.schemas import (
    AuditEventOut,
    BillingPeriodOut,
    DashboardOut,
    ExtendRequest,
    ExpenseCreate,
    ExpenseOut,
    ExpenseVoid,
    LoginRequest,
    NotificationOut,
    PageInfo,
    PaymentAdjustmentCreate,
    PaymentAdjustmentOut,
    PaymentCreate,
    PaymentOut,
    RefreshRequest,
    ResidentCreate,
    ResidentDetail,
    ResidentOut,
    ResidentUpdate,
    RoomCreate,
    RoomOut,
    RoomUpdate,
    SystemSettingsIn,
    SystemSettingsOut,
    TokenPair,
    UserCreate,
    UserOut,
    UserUpdate,
)
from app.security import create_token_pair, hash_password, verify_password
from app.services import audit_out, expense_out, notification_out, payment_out, resident_detail, resident_out, room_out, today_local
from app.residency import assert_room_available, close_open_period, record_room_assignment

router = APIRouter(prefix="/api")
UZ_MONTHS = ["yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul", "avgust", "sentabr", "oktabr", "noyabr", "dekabr"]
METHOD_UZ = {"CASH": "Naqd pul", "BANK": "Bank o‘tkazmasi", "CLICK": "Click", "PAYME": "Payme", "OTHER": "Boshqa"}
STATUS_UZ = {
    "PAID": "To‘langan", "CURRENT": "Amalda", "DUE_SOON": "Yaqinlashmoqda",
    "DUE_TODAY": "Bugun to‘lanadi", "OVERDUE": "Muddati o‘tgan",
    "NOT_CONFIGURED": "Davr belgilanmagan",
}
EXPENSE_CATEGORY_UZ = {
    "UTILITIES": "Kommunal xizmatlar", "REPAIR": "Ta'mirlash", "SUPPLIES": "Xo'jalik mollari",
    "FOOD": "Oziq-ovqat", "SALARY": "Ish haqi", "OTHER": "Boshqa",
}


def natural_room_key(number: str) -> tuple[tuple[int, object], ...]:
    return tuple((1, int(part)) if part.isdigit() else (0, part.casefold()) for part in re.split(r"(\d+)", number))


def page_result(items: list, total: int, skip: int, limit: int) -> PageInfo:
    return PageInfo(items=items, total=total, skip=skip, limit=limit)


def ensure_due_periods(db: Session, residents: list[Resident], actor: User) -> None:
    today = today_local()
    created = False
    for resident in residents:
        resident = db.scalar(select(Resident).where(Resident.id == resident.id).with_for_update())
        if not resident or resident.is_archived or current_period(db, resident.id):
            continue
        latest = db.scalar(
            select(BillingPeriod)
            .where(BillingPeriod.resident_id == resident.id, BillingPeriod.status != "CLOSED")
            .order_by(BillingPeriod.end_date.desc(), BillingPeriod.created_at.desc())
        )
        if latest and latest.status == "PAID" and latest.end_date <= today:
            new_period = ensure_period(db, resident, actor, as_of=today)
            created = created or bool(new_period and new_period.id != latest.id)
    if created:
        db.flush()


def verify_room_capacity(db: Session, room_id: UUID | None, resident_id: UUID | None = None, start_date: date | None = None, end_date: date | None = None) -> Room | None:
    if room_id is None:
        return None
    resident = db.get(Resident, resident_id) if resident_id else None
    return assert_room_available(
        db,
        room_id,
        start_date or (resident.check_in_date if resident and resident.check_in_date else today_local()),
        end_date,
        exclude_resident_id=resident_id,
    )

def net_payment(db: Session, payment: Payment) -> Decimal:
    adjustments = db.scalars(select(PaymentAdjustment).where(PaymentAdjustment.payment_id == payment.id)).all()
    return max(Decimal("0"), payment.amount + sum((item.amount_delta for item in adjustments), Decimal("0")))


def get_report_rows(
    db: Session,
    kind: str,
    date_from: date | None = None,
    date_to: date | None = None,
    interval: Literal["day", "week", "month"] = "month",
) -> tuple[list[str], list[list[object]]]:
    if kind == "residents":
        headers = ["Ism", "Familiya", "Telefon", "Telegram", "Xona", "Ko‘chib kirgan sana", "Davr (kun)", "Keyingi muddat", "Kelishilgan narx (so‘m)", "Qarzdorlik (so‘m)", "Holat"]
        residents = db.scalars(select(Resident).options(joinedload(Resident.room)).where(Resident.is_archived.is_(False)).order_by(Resident.last_name, Resident.first_name)).all()
        rows = []
        for resident in residents:
            item = resident_out(db, resident)
            rows.append([item.first_name, item.last_name, item.phone or "", item.telegram_username or "", item.room_number or "", item.check_in_date or "", item.payment_period_days or "", item.billing_end or "", item.agreed_price or 0, item.balance, STATUS_UZ[item.billing_status]])
        return headers, rows
    if kind == "expenses":
        start = date_from or date.min
        end = date_to or date.max
        expenses = db.scalars(
            select(Expense).where(Expense.expense_date.between(start, end), Expense.voided_at.is_(None))
            .order_by(Expense.expense_date.desc(), Expense.created_at.desc())
        ).all()
        headers = ["Sana", "Toifa", "Izoh", "Summa (so'm)", "Usul", "Hujjat raqami", "Qayd etdi"]
        rows = [[item.expense_date, EXPENSE_CATEGORY_UZ[item.category], item.description, item.amount, METHOD_UZ[item.method], item.reference, db.get(User, item.recorded_by).full_name if db.get(User, item.recorded_by) else ""] for item in expenses]
        return headers, rows
    if kind == "rooms":
        headers = ["Xona", "Qavat", "Sig‘im", "Band", "Bo‘sh o‘rin", "Narx (so‘m)", "Izoh"]
        rooms = sorted(db.scalars(select(Room).where(Room.is_active.is_(True))).all(), key=lambda room: natural_room_key(room.number))
        rows = []
        for room in rooms:
            item = room_out(db, room)
            rows.append([room.number, room.floor, room.capacity, item.resident_count, item.available_beds, room.price or "", room.notes])
        return headers, rows
    if kind in {"payments", "debts", "overdue", "revenue"}:
        start = date_from or date.min
        end = date_to or date.max
        payments = db.scalars(
            select(Payment).options(joinedload(Payment.resident)).where(Payment.payment_date.between(start, end)).order_by(Payment.payment_date.desc(), Payment.created_at.desc())
        ).all()
        if kind == "payments":
            headers = ["Yashovchi", "Summa (so‘m)", "Sana", "To‘lov usuli", "Davr boshlanishi", "Davr tugashi", "Qayd etdi", "Izoh"]
            rows = [[f"{p.resident.first_name} {p.resident.last_name}", net_payment(db, p), p.payment_date, METHOD_UZ[p.method], p.covered_start or "", p.covered_end or "", db.get(User, p.recorded_by).full_name if db.get(User, p.recorded_by) else "", p.notes] for p in payments]
            return headers, rows
        if kind == "revenue":
            headers = ["Davr boshlanishi", "Haqiqiy tushum (so‘m)"]
            buckets: dict[date, Decimal] = {}
            for payment in payments:
                if interval == "day":
                    bucket = payment.payment_date
                elif interval == "week":
                    bucket = payment.payment_date - timedelta(days=payment.payment_date.weekday())
                else:
                    bucket = payment.payment_date.replace(day=1)
                buckets[bucket] = buckets.get(bucket, Decimal("0")) + net_payment(db, payment)
            rows = [[bucket, amount] for bucket, amount in sorted(buckets.items())]
            return headers, rows
    if kind in {"debts", "overdue"}:
        headers = ["Yashovchi", "Xona", "Muddat", "Qarzdorlik (so‘m)", "Kechikish (kun)", "Holat"]
        residents = db.scalars(select(Resident).options(joinedload(Resident.room)).where(Resident.is_archived.is_(False)).order_by(Resident.last_name, Resident.first_name)).all()
        rows = []
        for resident in residents:
            item = resident_out(db, resident)
            if item.balance <= 0 or (kind == "overdue" and item.billing_status != "OVERDUE"):
                continue
            rows.append([f"{item.first_name} {item.last_name}", item.room_number or "", item.billing_end or "", item.balance, item.overdue_days, STATUS_UZ[item.billing_status]])
        return headers, rows
    raise HTTPException(status_code=404, detail="Hisobot turi topilmadi.")


@router.get("/health")
def health(db: Session = Depends(get_db)) -> dict[str, str]:
    db.execute(select(1))
    return {"status": "ok"}


@router.post("/auth/login", response_model=TokenPair)
def login(payload: LoginRequest, db: Session = Depends(get_db)) -> TokenPair:
    user = db.scalar(select(User).where(User.email == str(payload.email).lower()))
    if not user or not user.is_active or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Elektron pochta yoki parol noto‘g‘ri.")
    audit(db, user, "LOGIN", "User", user.id)
    db.commit()
    return TokenPair(**create_token_pair(user.id))


@router.post("/auth/refresh", response_model=TokenPair)
def refresh(payload: RefreshRequest, db: Session = Depends(get_db)) -> TokenPair:
    try:
        token_payload = jwt.decode(payload.refresh_token, settings.jwt_secret, algorithms=[settings.jwt_algorithm])
        if token_payload.get("type") != "refresh":
            raise ValueError("wrong token type")
        user_id = UUID(token_payload["sub"])
    except Exception:
        raise HTTPException(status_code=401, detail="Sessiyani yangilab bo‘lmadi. Qayta kiring.") from None
    user = db.get(User, user_id)
    if not user or not user.is_active:
        raise HTTPException(status_code=401, detail="Foydalanuvchi faol emas.")
    return TokenPair(**create_token_pair(user.id))


@router.get("/auth/me", response_model=UserOut)
def me(user: User = Depends(get_current_user)) -> User:
    return user


@router.get("/users", response_model=list[UserOut])
def list_users(db: Session = Depends(get_db), _: User = Depends(require_operator)) -> list[User]:
    return db.scalars(select(User).order_by(User.created_at)).all()


@router.post("/users", response_model=UserOut, status_code=201)
def create_user(payload: UserCreate, db: Session = Depends(get_db), actor: User = Depends(require_operator)) -> User:
    user = User(email=str(payload.email).lower(), full_name=payload.full_name.strip(), password_hash=hash_password(payload.password), role="ADMIN", is_active=True, is_primary=False)
    db.add(user)
    db.flush()
    audit(db, actor, "USER_CREATED", "User", user.id, {"email": user.email, "role": user.role})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Bu elektron pochta allaqachon ro‘yxatdan o‘tgan.") from None
    db.refresh(user)
    return user


@router.patch("/users/{user_id}", response_model=UserOut)
def update_user(user_id: UUID, payload: UserUpdate, db: Session = Depends(get_db), actor: User = Depends(require_operator)) -> User:
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Foydalanuvchi topilmadi.")
    values = payload.model_dump(exclude_unset=True)
    if user.is_primary and values.get("is_active") is False:
        raise HTTPException(status_code=409, detail="Asosiy operatorni faolsizlantirib bo‘lmaydi.")
    if "full_name" in values and values["full_name"] is not None:
        user.full_name = values["full_name"].strip()
    if values.get("password"):
        user.password_hash = hash_password(values["password"])
    if "is_active" in values and values["is_active"] is not None:
        user.is_active = values["is_active"]
    audit(db, actor, "USER_UPDATED", "User", user.id, {"changed_fields": sorted(k for k in values if k != "password")})
    db.commit()
    db.refresh(user)
    return user


@router.get("/rooms", response_model=PageInfo)
def list_rooms(
    status_filter: Literal["all", "available", "full"] = Query(default="all", alias="status"),
    skip: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db), _: User = Depends(get_current_user),
) -> PageInfo:
    rooms = sorted(db.scalars(select(Room).where(Room.is_active.is_(True))).all(), key=lambda room: natural_room_key(room.number))
    items = [room_out(db, room) for room in rooms]
    if status_filter == "available":
        items = [room for room in items if room.available_beds > 0]
    elif status_filter == "full":
        items = [room for room in items if room.available_beds == 0]
    return page_result(items[skip:skip + limit], len(items), skip, limit)


@router.post("/rooms", response_model=RoomOut, status_code=201)
def create_room(payload: RoomCreate, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> RoomOut:
    dormitory = db.get(Dormitory, payload.dormitory_id) if payload.dormitory_id else db.scalar(select(Dormitory).where(Dormitory.is_default.is_(True)))
    if payload.dormitory_id and dormitory is None:
        raise HTTPException(status_code=422, detail="Yotoqxona topilmadi.")
    if dormitory is None:
        dormitory = Dormitory(name="RS Best Hostel", is_default=True, is_active=True)
        db.add(dormitory)
        db.flush()
    if not dormitory.is_active:
        raise HTTPException(status_code=422, detail="Faol yotoqxonani tanlang.")
    room = Room(dormitory_id=dormitory.id, number=payload.number.strip(), floor=payload.floor, capacity=payload.capacity, price=payload.price, notes=payload.notes.strip())
    db.add(room)
    db.flush()
    audit(db, actor, "ROOM_CREATED", "Room", room.id, {"number": room.number})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Bu xona raqami mavjud.") from None
    db.refresh(room)
    return room_out(db, room)


@router.patch("/rooms/{room_id}", response_model=RoomOut)
def update_room(room_id: UUID, payload: RoomUpdate, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> RoomOut:
    room = db.get(Room, room_id)
    if not room or not room.is_active:
        raise HTTPException(status_code=404, detail="Xona topilmadi.")
    values = payload.model_dump(exclude_unset=True)
    if "dormitory_id" in values and values["dormitory_id"] != room.dormitory_id:
        target_dormitory = db.get(Dormitory, values["dormitory_id"]) if values["dormitory_id"] else None
        if not target_dormitory or not target_dormitory.is_active:
            raise HTTPException(status_code=422, detail="Faol yotoqxonani tanlang.")
        assigned_periods = db.scalar(select(func.count(ResidencePeriod.id)).where(ResidencePeriod.room_id == room.id)) or 0
        if assigned_periods:
            raise HTTPException(status_code=409, detail="Yashash tarixi bor xonaning yotoqxonasini almashtirib bo‘lmaydi.")
    occupant_count = db.scalar(select(func.count(Resident.id)).where(Resident.room_id == room.id, Resident.is_archived.is_(False))) or 0
    if values.get("capacity", room.capacity) < occupant_count:
        raise HTTPException(status_code=409, detail=f"Xonada {occupant_count} nafar yashovchi bor. Sig‘imni kamaytirib bo‘lmaydi.")
    for key, value in values.items():
        setattr(room, key, value.strip() if key in {"number", "notes"} and isinstance(value, str) else value)
    audit(db, actor, "ROOM_UPDATED", "Room", room.id, {"changed_fields": sorted(values)})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Bu xona raqami mavjud.") from None
    db.refresh(room)
    return room_out(db, room)


@router.delete("/rooms/{room_id}", status_code=204)
def archive_room(room_id: UUID, db: Session = Depends(get_db), actor: User = Depends(require_operator)) -> Response:
    room = db.get(Room, room_id)
    if not room or not room.is_active:
        raise HTTPException(status_code=404, detail="Xona topilmadi.")
    occupants = db.scalar(select(func.count(ResidencePeriod.id)).where(ResidencePeriod.room_id == room.id, ResidencePeriod.status.in_(("ACTIVE", "PLANNED")))) or 0
    if not occupants:
        occupants = db.scalar(select(func.count(Resident.id)).where(Resident.room_id == room.id, Resident.is_archived.is_(False))) or 0
    if occupants:
        raise HTTPException(status_code=409, detail="Yashovchilari bor xonani arxivlashdan oldin ularni boshqa xonaga o‘tkazing.")
    room.is_active = False
    audit(db, actor, "ROOM_ARCHIVED", "Room", room.id, {"number": room.number})
    db.commit()
    return Response(status_code=204)


@router.get("/residents", response_model=PageInfo)
def list_residents(
    q: str = "", room_id: UUID | None = None, payment_status: str = "all", archived: bool = False,
    sort: str = "last_name", direction: Literal["asc", "desc"] = "asc",
    skip: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db), actor: User = Depends(get_current_user),
) -> PageInfo:
    query = select(Resident).options(joinedload(Resident.room)).where(Resident.is_archived.is_(archived))
    if room_id:
        query = query.where(Resident.room_id == room_id)
    if q.strip():
        term = f"%{q.strip()}%"
        query = query.outerjoin(Room, Resident.room_id == Room.id).where(
            or_(Resident.first_name.ilike(term), Resident.last_name.ilike(term), Resident.phone.ilike(term), Resident.telegram_username.ilike(term), Room.number.ilike(term))
        )
    order_col = {"first_name": Resident.first_name, "created_at": Resident.created_at}.get(sort, Resident.last_name)
    query = query.order_by(order_col.desc() if direction == "desc" else order_col.asc())
    records = db.scalars(query).unique().all()
    items = [resident_out(db, resident) for resident in records]
    if payment_status != "all":
        items = [item for item in items if item.billing_status == payment_status.upper()]
    total = len(items)
    return page_result(items[skip:skip + limit], total, skip, limit)


@router.post("/residents", response_model=ResidentOut, status_code=201)
def create_resident(payload: ResidentCreate, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> ResidentOut:
    if payload.room_id:
        verify_room_capacity(db, payload.room_id, start_date=payload.check_in_date or today_local())
    resident = Resident(
        first_name=payload.first_name.strip(), last_name=payload.last_name.strip(),
        phone=payload.phone.strip() if payload.phone else None,
        telegram_username=payload.telegram_username.strip().removeprefix("@") if payload.telegram_username else None,
        room_id=payload.room_id, check_in_date=payload.check_in_date,
        payment_period_days=payload.payment_period_days, agreed_price=payload.agreed_price,
        notes=payload.notes.strip(),
    )
    db.add(resident)
    db.flush()
    if payload.room_id:
        record_room_assignment(db, resident, payload.room_id, payload.check_in_date or today_local(), None, actor)
    start = payload.billing_start or ((payload.check_in_date or today_local()) if payload.payment_period_days and payload.agreed_price else None)
    end = payload.billing_end
    days = payload.payment_period_days
    if start and days and not end:
        end = start + timedelta(days=days)
    if start and end:
        if not days:
            days = (end - start).days
            resident.payment_period_days = days
        if resident.agreed_price is None:
            raise HTTPException(status_code=422, detail="To‘lov davrini belgilash uchun kelishilgan narxni kiriting.")
        create_initial_period(db, resident, start, end, actor)
    audit(db, actor, "RESIDENT_CREATED", "Resident", resident.id, {"name": f"{resident.first_name} {resident.last_name}"})
    notify(db, "Yangi yashovchi", f"{resident.first_name} {resident.last_name} ro‘yxatga qo‘shildi.", "RESIDENT", resident.id, actor.id)
    db.commit()
    db.refresh(resident)
    return resident_out(db, resident)


@router.get("/residents/{resident_id}", response_model=ResidentDetail)
def get_resident(resident_id: UUID, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> ResidentDetail:
    resident = db.scalar(select(Resident).options(joinedload(Resident.room)).where(Resident.id == resident_id))
    if not resident:
        raise HTTPException(status_code=404, detail="Yashovchi topilmadi.")
    return resident_detail(db, resident)


@router.patch("/residents/{resident_id}", response_model=ResidentOut)
def update_resident(resident_id: UUID, payload: ResidentUpdate, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> ResidentOut:
    resident = db.get(Resident, resident_id)
    if not resident:
        raise HTTPException(status_code=404, detail="Yashovchi topilmadi.")
    values = payload.model_dump(exclude_unset=True)
    billing_keys = {"billing_start", "billing_end", "payment_period_days", "agreed_price"}
    if billing_keys.intersection(values) and actor.role != "OPERATOR" and not settings.admins_can_adjust_billing:
        raise HTTPException(status_code=403, detail="To‘lov davrini faqat operator o‘zgartira oladi.")
    next_room = values.get("room_id", resident.room_id)
    next_start = values.get("check_in_date", resident.check_in_date)
    if "room_id" in values or "check_in_date" in values:
        if next_room:
            if "check_in_date" not in values and next_room != resident.room_id:
                next_start = today_local()
            verify_room_capacity(db, next_room, resident.id, next_start or today_local())
    old_billing = current_period(db, resident.id)
    old_start = old_billing.start_date if old_billing else None
    old_end = old_billing.end_date if old_billing else None
    old_days = resident.payment_period_days
    for key, value in values.items():
        if key in {"first_name", "last_name", "notes"} and isinstance(value, str):
            value = value.strip()
        if key == "phone" and isinstance(value, str):
            value = value.strip() or None
        if key == "telegram_username" and isinstance(value, str):
            value = value.strip().removeprefix("@") or None
        setattr(resident, key, value)
    if "room_id" in values or "check_in_date" in values:
        if next_room:
            record_room_assignment(db, resident, next_room, next_start or today_local(), None, actor)
        else:
            close_open_period(db, resident)
    billing_changed = bool(billing_keys.intersection(values))
    if billing_changed:
        start = values.get("billing_start", old_start)
        days = values.get("payment_period_days", old_days)
        if start is None and old_billing is None and days and resident.agreed_price and resident.agreed_price > 0:
            start = resident.check_in_date or today_local()
        if "billing_end" in values:
            end = values["billing_end"]
        elif "billing_start" in values or "payment_period_days" in values:
            end = start + timedelta(days=days) if start is not None and days else old_end
        else:
            end = old_end
        if start is not None and end is not None:
            days = days or (end - start).days
            set_period_dates(db, resident, actor, start, end, days)
        elif old_billing:
            raise HTTPException(status_code=409, detail="To‘lov qaydi bor davrni o‘chirib bo‘lmaydi. Yangi sanalarni belgilang.")
    audit(db, actor, "RESIDENT_UPDATED", "Resident", resident.id, {"changed_fields": sorted(values)})
    db.commit()
    db.refresh(resident)
    ensure_due_periods(db, [resident], actor)
    db.commit()
    return resident_out(db, resident)


@router.post("/residents/{resident_id}/archive", response_model=ResidentOut)
def archive_resident(resident_id: UUID, db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> ResidentOut:
    resident = db.get(Resident, resident_id)
    if not resident:
        raise HTTPException(status_code=404, detail="Yashovchi topilmadi.")
    if not resident.is_archived:
        resident.is_archived = True
        resident.archived_at = datetime.now(UTC)
        close_open_period(db, resident, resident.archived_at)
        if settings.file_retention_days:
            application_ids = select(StudentApplication.id).where(StudentApplication.resident_id == resident.id)
            attachments = db.scalars(select(ApplicationFile).where(ApplicationFile.application_id.in_(application_ids), ApplicationFile.deleted_at.is_(None))).all()
            for attachment in attachments:
                attachment.delete_after = resident.archived_at + timedelta(days=settings.file_retention_days)
        audit(db, actor, "RESIDENT_ARCHIVED", "Resident", resident.id)
        db.commit()
    return resident_out(db, resident)


@router.post("/residents/{resident_id}/restore", response_model=ResidentOut)
def restore_resident(resident_id: UUID, db: Session = Depends(get_db), actor: User = Depends(require_operator)) -> ResidentOut:
    resident = db.get(Resident, resident_id)
    if not resident:
        raise HTTPException(status_code=404, detail="Yashovchi topilmadi.")
    verify_room_capacity(db, resident.room_id, resident.id, today_local())
    resident.is_archived = False
    resident.archived_at = None
    if resident.room_id:
        record_room_assignment(db, resident, resident.room_id, today_local(), None, actor)
    audit(db, actor, "RESIDENT_RESTORED", "Resident", resident.id)
    db.commit()
    ensure_due_periods(db, [resident], actor)
    db.commit()
    return resident_out(db, resident)


@router.post("/residents/{resident_id}/extend", response_model=BillingPeriodOut)
def extend_resident(resident_id: UUID, payload: ExtendRequest, db: Session = Depends(get_db), actor: User = Depends(require_billing_permission)) -> BillingPeriodOut:
    resident = db.get(Resident, resident_id)
    if not resident:
        raise HTTPException(status_code=404, detail="Yashovchi topilmadi.")
    period = extend_period(db, resident, actor, payload.days, payload.end_date, payload.reason)
    notify(db, "Yashash muddati uzaytirildi", f"{resident.first_name} {resident.last_name} uchun muddat yangilandi. To‘lov alohida qayd etiladi.", "BILLING", resident.id, actor.id)
    db.commit()
    db.refresh(period)
    return BillingPeriodOut(**{
        "id": period.id, "start_date": period.start_date, "end_date": period.end_date,
        "period_days": period.period_days, "charge_amount": period.charge_amount,
        "paid_amount": period_paid(db, period.id), "balance": period_balance(db, period), "status": period.status,
    })


@router.get("/residents/{resident_id}/payments", response_model=list[PaymentOut])
def resident_payments(resident_id: UUID, db: Session = Depends(get_db), _: User = Depends(get_current_user)) -> list[PaymentOut]:
    if not db.get(Resident, resident_id):
        raise HTTPException(status_code=404, detail="Yashovchi topilmadi.")
    rows = db.scalars(select(Payment).options(joinedload(Payment.resident)).where(Payment.resident_id == resident_id).order_by(Payment.payment_date.desc(), Payment.created_at.desc())).all()
    return [payment_out(db, row) for row in rows]


@router.get("/payments", response_model=PageInfo)
def list_payments(
    resident_id: UUID | None = None, date_from: date | None = None, date_to: date | None = None,
    skip: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db), _: User = Depends(get_current_user),
) -> PageInfo:
    query = select(Payment).options(joinedload(Payment.resident)).order_by(Payment.payment_date.desc(), Payment.created_at.desc())
    if resident_id:
        query = query.where(Payment.resident_id == resident_id)
    if date_from:
        query = query.where(Payment.payment_date >= date_from)
    if date_to:
        query = query.where(Payment.payment_date <= date_to)
    rows = db.scalars(query).unique().all()
    total = len(rows)
    return page_result([payment_out(db, item) for item in rows[skip:skip + limit]], total, skip, limit)


@router.post("/payments", response_model=PaymentOut, status_code=201)
def create_payment(
    payload: PaymentCreate,
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key", max_length=128),
    db: Session = Depends(get_db), actor: User = Depends(get_current_user),
) -> PaymentOut:
    if idempotency_key is not None and (len(idempotency_key) < 8 or any(not (c.isalnum() or c in "._:-") for c in idempotency_key)):
        raise HTTPException(status_code=422, detail="Idempotency-Key kamida 8 ta xavfsiz belgidan iborat bo'lishi kerak.")
    if idempotency_key:
        previous = db.scalar(select(Payment).options(joinedload(Payment.resident)).where(Payment.recorded_by == actor.id, Payment.idempotency_key == idempotency_key))
        if previous:
            same_request = (
                previous.resident_id == payload.resident_id
                and previous.amount == payload.amount
                and previous.payment_date == payload.payment_date
                and previous.method == payload.method
                and previous.notes == payload.notes.strip()
                and (payload.billing_period_id is None or previous.billing_period_id == payload.billing_period_id)
            )
            if not same_request:
                raise HTTPException(status_code=409, detail="Bu kalit boshqa to'lov uchun ishlatilgan.")
            return payment_out(db, previous)
    if payload.payment_date > today_local():
        raise HTTPException(status_code=422, detail="To'lov sanasi kelajakda bo'lishi mumkin emas.")
    resident = db.scalar(select(Resident).where(Resident.id == payload.resident_id).with_for_update())
    if not resident or resident.is_archived:
        raise HTTPException(status_code=404, detail="Faol yashovchi topilmadi.")
    if idempotency_key:
        previous = db.scalar(select(Payment).options(joinedload(Payment.resident)).where(Payment.recorded_by == actor.id, Payment.idempotency_key == idempotency_key))
        if previous:
            if previous.resident_id == payload.resident_id and previous.amount == payload.amount and previous.payment_date == payload.payment_date and previous.method == payload.method and previous.notes == payload.notes.strip() and (payload.billing_period_id is None or previous.billing_period_id == payload.billing_period_id):
                return payment_out(db, previous)
            raise HTTPException(status_code=409, detail="Bu kalit boshqa to'lov uchun ishlatilgan.")
    ensure_due_periods(db, [resident], actor)
    if payload.billing_period_id:
        period_query = select(BillingPeriod).where(BillingPeriod.id == payload.billing_period_id, BillingPeriod.resident_id == resident.id)
    else:
        period_query = select(BillingPeriod).where(BillingPeriod.resident_id == resident.id, BillingPeriod.status == "OPEN").order_by(BillingPeriod.start_date.desc())
    period = db.scalar(period_query.with_for_update())
    if not period or period.status != "OPEN":
        raise HTTPException(status_code=409, detail="Yashovchi uchun ochiq to'lov davri topilmadi.")
    balance = period_balance(db, period)
    if balance <= 0:
        raise HTTPException(status_code=409, detail="Bu davr uchun qarzdorlik qolmagan.")
    if payload.amount > balance:
        raise HTTPException(status_code=422, detail="To'lov summasi davr qoldig'idan oshmasligi kerak.")
    payment = Payment(
        resident_id=resident.id, billing_period_id=period.id, amount=payload.amount,
        payment_date=payload.payment_date, method=payload.method,
        covered_start=period.start_date, covered_end=period.end_date,
        notes=payload.notes.strip(), idempotency_key=idempotency_key, recorded_by=actor.id,
    )
    db.add(payment)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        if idempotency_key:
            previous = db.scalar(select(Payment).options(joinedload(Payment.resident)).where(Payment.recorded_by == actor.id, Payment.idempotency_key == idempotency_key))
            if previous and previous.resident_id == payload.resident_id and previous.amount == payload.amount and previous.payment_date == payload.payment_date and previous.method == payload.method and previous.notes == payload.notes.strip() and (payload.billing_period_id is None or previous.billing_period_id == payload.billing_period_id):
                return payment_out(db, previous)
        raise HTTPException(status_code=409, detail="To'lov qayta yuborilgan. Ma'lumotni yangilang.") from None
    audit(db, actor, "PAYMENT_RECORDED", "Payment", payment.id, {"resident_id": str(resident.id), "amount": str(payment.amount), "method": payment.method})
    notify(db, "Yangi to'lov qayd etildi", f"{resident.first_name} {resident.last_name} uchun {payment.amount} so'm qabul qilindi.", "PAYMENT", resident.id, actor.id)
    mark_paid_and_advance(db, period, actor)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        if idempotency_key:
            previous = db.scalar(select(Payment).options(joinedload(Payment.resident)).where(Payment.recorded_by == actor.id, Payment.idempotency_key == idempotency_key))
            if previous and previous.resident_id == payload.resident_id and previous.amount == payload.amount and previous.payment_date == payload.payment_date and previous.method == payload.method and previous.notes == payload.notes.strip():
                return payment_out(db, previous)
        raise HTTPException(status_code=409, detail="To'lov qayta yuborilgan yoki davr qoldig'i o'zgargan. Ma'lumotni yangilang.") from None
    db.refresh(payment)
    return payment_out(db, payment)


@router.get("/expenses", response_model=PageInfo)
def list_expenses(
    date_from: date | None = None, date_to: date | None = None,
    include_voided: bool = False, skip: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db), _: User = Depends(get_current_user),
) -> PageInfo:
    if date_from and date_to and date_to < date_from:
        raise HTTPException(status_code=422, detail="Yakuniy sana boshlang'ich sanadan oldin bo'lishi mumkin emas.")
    query = select(Expense).order_by(Expense.expense_date.desc(), Expense.created_at.desc())
    if date_from:
        query = query.where(Expense.expense_date >= date_from)
    if date_to:
        query = query.where(Expense.expense_date <= date_to)
    if not include_voided:
        query = query.where(Expense.voided_at.is_(None))
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.scalars(query.offset(skip).limit(limit)).all()
    return page_result([expense_out(db, row) for row in rows], total, skip, limit)


@router.post("/expenses", response_model=ExpenseOut, status_code=201)
def create_expense(
    payload: ExpenseCreate,
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key", max_length=128),
    db: Session = Depends(get_db), actor: User = Depends(get_current_user),
) -> ExpenseOut:
    if idempotency_key is not None and (len(idempotency_key) < 8 or any(not (c.isalnum() or c in "._:-") for c in idempotency_key)):
        raise HTTPException(status_code=422, detail="Idempotency-Key kamida 8 ta xavfsiz belgidan iborat bo'lishi kerak.")
    if idempotency_key:
        previous = db.scalar(select(Expense).where(Expense.recorded_by == actor.id, Expense.idempotency_key == idempotency_key))
        if previous:
            same_request = previous.expense_date == payload.expense_date and previous.category == payload.category and previous.description == payload.description.strip() and previous.amount == payload.amount and previous.method == payload.method and previous.reference == payload.reference.strip()
            if not same_request:
                raise HTTPException(status_code=409, detail="Bu kalit boshqa xarajat uchun ishlatilgan.")
            return expense_out(db, previous)
    if payload.expense_date > today_local():
        raise HTTPException(status_code=422, detail="Xarajat sanasi kelajakda bo'lishi mumkin emas.")
    expense = Expense(
        expense_date=payload.expense_date, category=payload.category,
        description=payload.description.strip(), amount=payload.amount,
        method=payload.method, reference=payload.reference.strip(),
        idempotency_key=idempotency_key, recorded_by=actor.id,
    )
    db.add(expense)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        if idempotency_key:
            previous = db.scalar(select(Expense).where(Expense.recorded_by == actor.id, Expense.idempotency_key == idempotency_key))
            if previous and previous.expense_date == payload.expense_date and previous.category == payload.category and previous.description == payload.description.strip() and previous.amount == payload.amount and previous.method == payload.method and previous.reference == payload.reference.strip():
                return expense_out(db, previous)
        raise HTTPException(status_code=409, detail="Xarajat qayta yuborilgan. Ma'lumotni yangilang.") from None
    audit(db, actor, "EXPENSE_RECORDED", "Expense", expense.id, {"category": expense.category, "amount": str(expense.amount), "method": expense.method})
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        if idempotency_key:
            previous = db.scalar(select(Expense).where(Expense.recorded_by == actor.id, Expense.idempotency_key == idempotency_key))
            if previous and previous.expense_date == payload.expense_date and previous.category == payload.category and previous.description == payload.description.strip() and previous.amount == payload.amount and previous.method == payload.method and previous.reference == payload.reference.strip():
                return expense_out(db, previous)
        raise HTTPException(status_code=409, detail="Xarajat qayta yuborilgan. Ma'lumotni yangilang.") from None
    db.refresh(expense)
    return expense_out(db, expense)


@router.post("/expenses/{expense_id}/void", response_model=ExpenseOut)
def void_expense(expense_id: UUID, payload: ExpenseVoid, db: Session = Depends(get_db), actor: User = Depends(require_operator)) -> ExpenseOut:
    expense = db.scalar(select(Expense).where(Expense.id == expense_id).with_for_update())
    if not expense:
        raise HTTPException(status_code=404, detail="Xarajat topilmadi.")
    if expense.voided_at:
        raise HTTPException(status_code=409, detail="Xarajat avval bekor qilingan.")
    expense.voided_at = datetime.now(UTC)
    expense.voided_by = actor.id
    expense.void_reason = payload.reason.strip()
    audit(db, actor, "EXPENSE_VOIDED", "Expense", expense.id, {"amount": str(expense.amount), "reason": expense.void_reason})
    db.commit()
    db.refresh(expense)
    return expense_out(db, expense)


@router.post("/payments/{payment_id}/adjustments", response_model=PaymentAdjustmentOut, status_code=201)
def adjust_payment(payment_id: UUID, payload: PaymentAdjustmentCreate, db: Session = Depends(get_db), actor: User = Depends(require_operator)) -> PaymentAdjustmentOut:
    payment = db.scalar(select(Payment).options(joinedload(Payment.adjustments)).where(Payment.id == payment_id).with_for_update())
    if not payment:
        raise HTTPException(status_code=404, detail="To‘lov topilmadi.")
    current_net = payment.amount + sum((item.amount_delta for item in payment.adjustments), Decimal("0"))
    if current_net + payload.amount_delta < 0:
        raise HTTPException(status_code=422, detail="Tuzatish to‘lov qoldig‘ini manfiy qilmasligi kerak.")
    adjustment = PaymentAdjustment(payment_id=payment.id, amount_delta=payload.amount_delta, reason=payload.reason.strip(), adjusted_by=actor.id)
    db.add(adjustment)
    db.flush()
    period = db.get(BillingPeriod, payment.billing_period_id) if payment.billing_period_id else None
    if period:
        paid = period_paid(db, period.id)
        if paid < period.charge_amount:
            period.status = "OPEN"
            newer_open = db.scalars(select(BillingPeriod).where(BillingPeriod.resident_id == period.resident_id, BillingPeriod.status == "OPEN", BillingPeriod.start_date > period.start_date)).all()
            for later in newer_open:
                later.status = "CLOSED"
        else:
            period.status = "PAID"
            mark_paid_and_advance(db, period, actor)
    audit(db, actor, "PAYMENT_ADJUSTED", "Payment", payment.id, {"adjustment_id": str(adjustment.id), "amount_delta": str(payload.amount_delta), "reason": payload.reason.strip()})
    db.commit()
    db.refresh(adjustment)
    return adjustment


@router.get("/payments/{payment_id}/adjustments", response_model=list[PaymentAdjustmentOut])
def list_adjustments(payment_id: UUID, db: Session = Depends(get_db), _: User = Depends(require_operator)) -> list[PaymentAdjustment]:
    if not db.get(Payment, payment_id):
        raise HTTPException(status_code=404, detail="To‘lov topilmadi.")
    return db.scalars(select(PaymentAdjustment).where(PaymentAdjustment.payment_id == payment_id).order_by(PaymentAdjustment.created_at)).all()


@router.get("/dashboard", response_model=DashboardOut)
def dashboard(db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> DashboardOut:
    today = today_local()
    residents = db.scalars(select(Resident).options(joinedload(Resident.room)).where(Resident.is_archived.is_(False))).all()
    active_rooms = sorted(db.scalars(select(Room).where(Room.is_active.is_(True))).all(), key=lambda room: natural_room_key(room.number))
    items = [resident_out(db, item, today) for item in residents]
    capacity = sum(room.capacity for room in active_rooms)
    room_stats = [room_out(db, room) for room in active_rooms]
    assigned = sum(item.resident_count for item in room_stats)
    month_start = today.replace(day=1)
    first_trend_index = month_start.year * 12 + month_start.month - 1 - 5
    first_trend_month = date(first_trend_index // 12, first_trend_index % 12 + 1, 1)
    payments = db.scalars(select(Payment).options(joinedload(Payment.resident)).where(Payment.payment_date >= first_trend_month).order_by(Payment.payment_date)).all()
    month_income = sum((net_payment(db, item) for item in payments if item.payment_date >= month_start), Decimal("0"))
    expenses = db.scalars(select(Expense).where(Expense.expense_date >= month_start, Expense.voided_at.is_(None))).all()
    month_expenses = sum((item.amount for item in expenses), Decimal("0"))
    open_periods = db.scalars(
        select(BillingPeriod).join(Resident, BillingPeriod.resident_id == Resident.id).where(
            BillingPeriod.status == "OPEN", Resident.is_archived.is_(False)
        )
    ).all()
    expected = sum((item.charge_amount for item in open_periods), Decimal("0"))
    outstanding = sum((period_balance(db, item) for item in open_periods), Decimal("0"))
    trend = []
    for offset in range(5, -1, -1):
        month_index = month_start.year * 12 + month_start.month - 1 - offset
        month = date(month_index // 12, month_index % 12 + 1, 1)
        next_month = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
        amount = sum((net_payment(db, item) for item in payments if month <= item.payment_date < next_month), Decimal("0"))
        trend.append({"month": f"{UZ_MONTHS[month.month - 1]} {month.year}", "amount": amount})
    recent = db.scalars(select(Payment).options(joinedload(Payment.resident)).order_by(Payment.payment_date.desc(), Payment.created_at.desc()).limit(6)).all()
    new_residents = db.scalars(select(Resident).options(joinedload(Resident.room)).where(Resident.is_archived.is_(False)).order_by(Resident.created_at.desc()).limit(5)).all()
    return DashboardOut(
        residents=len(residents), rooms=len(active_rooms), occupancy_rate=round(assigned / capacity * 100, 1) if capacity else 0,
        available_beds=max(0, capacity - assigned), due_today=sum(x.billing_status == "DUE_TODAY" for x in items),
        due_soon=sum(x.billing_status == "DUE_SOON" for x in items), overdue=sum(x.billing_status == "OVERDUE" for x in items),
        actual_revenue=month_income, actual_expenses=month_expenses, net_cash_flow=month_income - month_expenses,
        expected_payments=expected, outstanding_debt=outstanding,
        revenue_trend=trend,
        occupancy_by_room=[{"room": item.number, "occupied": item.resident_count, "available": item.available_beds} for item in room_stats],
        recent_payments=[payment_out(db, item) for item in recent],
        recent_residents=[resident_out(db, item, today) for item in new_residents],
    )


@router.get("/notifications", response_model=list[NotificationOut])
def list_notifications(db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> list[NotificationOut]:
    stored = db.scalars(
        select(Notification).where(Notification.user_id == actor.id, Notification.read_at.is_(None))
        .order_by(Notification.created_at.desc()).limit(50)
    ).all()
    return [notification_out(item) for item in stored]


@router.post("/notifications/read-all", status_code=204)
def mark_notifications_read(db: Session = Depends(get_db), actor: User = Depends(get_current_user)) -> Response:
    rows = db.scalars(select(Notification).where(Notification.user_id == actor.id, Notification.read_at.is_(None))).all()
    now = datetime.now(UTC)
    for row in rows:
        row.read_at = now
    db.commit()
    return Response(status_code=204)


@router.get("/settings", response_model=SystemSettingsOut)
def get_settings(db: Session = Depends(get_db), _: User = Depends(get_current_user)) -> SystemSettingsOut:
    record = db.get(SystemSetting, "general")
    value = record.value if record else {"dorm_name": "RS Best Hostel", "default_period_days": 30, "currency": "UZS"}
    return SystemSettingsOut(**value, updated_at=record.updated_at if record else None, billing_change_allowed=settings.admins_can_adjust_billing)


@router.put("/settings", response_model=SystemSettingsOut)
def update_settings(payload: SystemSettingsIn, db: Session = Depends(get_db), actor: User = Depends(require_operator)) -> SystemSettingsOut:
    record = db.get(SystemSetting, "general")
    if not record:
        record = SystemSetting(key="general", value=payload.model_dump(), updated_by=actor.id)
        db.add(record)
    else:
        record.value = payload.model_dump()
        record.updated_by = actor.id
    audit(db, actor, "SETTINGS_UPDATED", "SystemSetting", details={"dorm_name": payload.dorm_name})
    db.commit()
    db.refresh(record)
    return SystemSettingsOut(**record.value, updated_at=record.updated_at, billing_change_allowed=settings.admins_can_adjust_billing)


@router.get("/audit", response_model=PageInfo)
def list_audit(
    skip: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db), _: User = Depends(require_operator),
) -> PageInfo:
    total = db.scalar(select(func.count(AuditEvent.id))) or 0
    rows = db.scalars(select(AuditEvent).order_by(AuditEvent.created_at.desc()).offset(skip).limit(limit)).all()
    return page_result([audit_out(db, item) for item in rows], total, skip, limit)


def spreadsheet_safe(value: object) -> object:
    if isinstance(value, str) and (value[:1] in ("\t", "\r", "\n") or value.lstrip()[:1] in ("=", "+", "-", "@")):
        return "'" + value
    return value


def xlsx_file(headers: list[str], rows: list[list[object]], sheet_name: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="255744")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([str(value) if isinstance(value, (date, datetime)) else spreadsheet_safe(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(42, max(12, max((len(str(cell.value or "")) for cell in column), default=10) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_file(headers: list[str], rows: list[list[object]], title: str) -> bytes:
    out = io.BytesIO()
    document = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=26, rightMargin=26, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    font_paths = ["C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    try:
        font_path = next(path for path in font_paths if exists(path))
        pdfmetrics.registerFont(TTFont("RS Best Hostel", font_path))
        styles["Normal"].fontName = "RS Best Hostel"
        styles["Title"].fontName = "RS Best Hostel"
    except Exception:
        pass
    content = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    data = [[str(item) for item in headers]]
    data.extend([[str(item) for item in row] for row in rows[:2500]])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#255744")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), "RS Best Hostel" if "RS Best Hostel" in pdfmetrics.getRegisteredFontNames() else "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDE4DD")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7F3")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    content.append(table)
    document.build(content)
    return out.getvalue()


@router.get("/reports/{kind}.{file_format}")
def export_report(
    kind: Literal["residents", "rooms", "payments", "debts", "overdue", "revenue", "expenses"],
    file_format: Literal["csv", "xlsx", "pdf"],
    date_from: date | None = None, date_to: date | None = None,
    interval: Literal["day", "week", "month"] = "month",
    db: Session = Depends(get_db), actor: User = Depends(get_current_user),
) -> Response:
    if date_from and date_to and date_to < date_from:
        raise HTTPException(status_code=422, detail="Yakuniy sana boshlang‘ich sanadan oldin bo‘lishi mumkin emas.")
    headers, rows = get_report_rows(db, kind, date_from, date_to, interval)
    name = f"rs-best-hostel-{kind}-{today_local().isoformat()}"
    if file_format == "csv":
        text_stream = io.StringIO(newline="")
        writer = csv.writer(text_stream)
        writer.writerow(headers)
        writer.writerows([[spreadsheet_safe(value) for value in row] for row in rows])
        content = ("\ufeff" + text_stream.getvalue()).encode("utf-8")
        return Response(content, media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})
    if file_format == "xlsx":
        sheet_names = {"residents": "Yashovchilar", "rooms": "Xonalar", "payments": "To‘lovlar", "debts": "Qarzdorlik", "overdue": "Muddati o‘tgan", "revenue": "Tushum", "expenses": "Xarajatlar"}
        content = xlsx_file(headers, rows, sheet_names[kind])
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        pdf_titles = {"residents": "Yashovchilar", "rooms": "Xonalar bandligi", "payments": "To‘lovlar tarixi", "debts": "Qarzdorliklar", "overdue": "Muddati o‘tgan to‘lovlar", "revenue": "Haqiqiy tushum", "expenses": "Xarajatlar"}
        content = pdf_file(headers, rows, f"RS Best Hostel — {pdf_titles[kind]}")
        media = "application/pdf"
    extension = "xlsx" if file_format == "xlsx" else "pdf"
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="{name}.{extension}"'})
